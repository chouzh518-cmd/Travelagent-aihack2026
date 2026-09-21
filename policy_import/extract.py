from __future__ import annotations

import re
from io import BytesIO

from .models import Block, Chunk, ExtractionOptions, Issue, Location, Reference, Table

PARSER_VERSION = "1.2"
ARTICLE = re.compile(r"^(第[0-9０-９一二三四五六七八九十百]+条)")
CHAPTER = re.compile(r"^第[0-9０-９一二三四五六七八九十百]+章")
TABLE_LABEL = re.compile(r"^[（(]?(別表[0-9０-９一二三四五六七八九十百]+)[）)]?")
REFERENCE = re.compile(r"別表[0-9０-９一二三四五六七八九十百]+|第[0-9０-９一二三四五六七八九十百]+条|前条|前項")
PLACEHOLDER = re.compile(r"○{2,}|○,|YYYY|MM月|DD日")
OCR_IMAGE_TYPES = {"png", "jpg", "jpeg", "tif", "tiff", "bmp", "webp"}


def recognize_image(image):
    """Recognize Japanese, Chinese, and Latin text locally with RapidOCR."""
    from rapidocr import EngineType, LangDet, LangRec, ModelType, OCRVersion, RapidOCR

    engine = getattr(recognize_image, "engine", None)
    if engine is None:
        engine = RapidOCR(params={
            "Det.engine_type": EngineType.ONNXRUNTIME, "Det.lang_type": LangDet.CH,
            "Det.model_type": ModelType.SMALL, "Det.ocr_version": OCRVersion.PPOCRV6,
            "Rec.engine_type": EngineType.ONNXRUNTIME, "Rec.lang_type": LangRec.JAPAN,
            "Rec.model_type": ModelType.SMALL, "Rec.ocr_version": OCRVersion.PPOCRV6,
        })
        recognize_image.engine = engine
    import numpy as np
    result = engine(np.asarray(image.convert("RGB")))
    lines = getattr(result, "txts", None) or []
    scores = getattr(result, "scores", None) or []
    return [(str(line).strip(), float(scores[i]) if i < len(scores) else None)
            for i, line in enumerate(lines) if str(line).strip()]


class ExtractionError(ValueError):
    pass


def pdf_blocks(data: bytes):
    import pdfplumber
    import pymupdf

    blocks, issues = [], []
    rendered_pdf = pymupdf.open(stream=data, filetype="pdf")
    if len(rendered_pdf) > 200:
        rendered_pdf.close()
        raise ExtractionError("PDF exceeds the 200 page limit")
    with pdfplumber.open(BytesIO(data)) as pdf:
        for page_number, page in enumerate(pdf.pages, 1):
            # Keep table cells as grids and prevent duplicate flattened text.
            tables = page.find_tables()
            events = []
            for table in tables:
                rows = table.extract()
                if not rows:
                    continue
                grid = [[list(map(float, cell)) if cell else None for cell in row.cells]
                        for row in table.rows]
                events.append((float(table.bbox[1]), "table", rows, list(map(float, table.bbox)), grid))
            def outside_tables(obj):
                if obj.get("object_type") != "char":
                    return True
                x = (obj["x0"] + obj["x1"]) / 2
                y = (obj["top"] + obj["bottom"]) / 2
                return not any(t.bbox[0] <= x <= t.bbox[2] and t.bbox[1] <= y <= t.bbox[3] for t in tables)
            for line in page.filter(outside_tables).extract_text_lines():
                if line["text"].strip():
                    events.append((float(line["top"]), "text", line["text"],
                                   [float(line[k]) for k in ("x0", "top", "x1", "bottom")], None))
            page_text = "\n".join(value for _, kind, value, _, _ in events if kind == "text")
            used_full_page_ocr = False
            if not page_text.strip() or (page.images and len(page_text.strip()) < 80):
                try:
                    rendered_page = rendered_pdf[page_number - 1]
                    if rendered_page.rect.width * rendered_page.rect.height * 4 > 50_000_000:
                        raise ExtractionError(f"PDF page {page_number} exceeds the 50 megapixel OCR limit")
                    pixmap = rendered_page.get_pixmap(matrix=pymupdf.Matrix(2, 2), alpha=False)
                    from PIL import Image
                    recognized = recognize_image(Image.frombytes("RGB", [pixmap.width, pixmap.height], pixmap.samples))
                    if recognized:
                        # Keep extracted table grids; OCR replaces sparse flattened text to avoid duplicates.
                        events = [event for event in events if event[1] == "table"]
                        used_full_page_ocr = True
                        text = "\n".join(line for line, _ in recognized)
                        events.append((0.0, "text", text, [0.0, 0.0, float(page.width), float(page.height)], None))
                        low_confidence = [score for _, score in recognized if score is not None and score < 0.55]
                        if low_confidence:
                            issues.append(Issue(code="OCR_LOW_CONFIDENCE", message="OCR 識別精度が低い箇所があります。原文画像と照合してください。", page_number=page_number))
                    else:
                        issues.append(Issue(code="OCR_NO_TEXT", message="このページから文字を認識できませんでした。原文を確認してください。", page_number=page_number))
                except Exception as exc:
                    raise ExtractionError(f"PDF page {page_number} OCR failed: {type(exc).__name__}: {exc}") from exc
            if not events:
                issues.append(Issue(code="PAGE_WITHOUT_TEXT", message="抽出可能な文字がありません。ページを確認してください。", page_number=page_number))
            elif page.images and not used_full_page_ocr:
                # OCR embedded photographs/screenshots that coexist with selectable page text.
                rendered_page = rendered_pdf[page_number - 1]
                from PIL import Image
                for image_info in rendered_page.get_images(full=True):
                    image_data = rendered_pdf.extract_image(image_info[0])
                    with Image.open(BytesIO(image_data["image"])) as embedded:
                        if embedded.width * embedded.height > 50_000_000:
                            continue
                        recognized = recognize_image(embedded.convert("RGB"))
                    for line, score in recognized:
                        rects = rendered_page.get_image_rects(image_info[0])
                        bbox = list(map(float, rects[0])) if rects else [0.0, 0.0, float(page.width), float(page.height)]
                        events.append((bbox[1], "text", line, bbox, None))
                        if score is not None and score < 0.55:
                            issues.append(Issue(code="OCR_LOW_CONFIDENCE", message="OCR 識別精度が低い箇所があります。原文画像と照合してください。", page_number=page_number))
            for _, kind, value, bbox, grid in sorted(events, key=lambda event: event[0]):
                location = Location(page_number=page_number, block_index=len(blocks)+1,
                                    locator=f"PDF page {page_number}, bbox {bbox}", bbox=bbox)
                table = Table(rows=value, location=location, cell_bboxes=grid) if kind == "table" else None
                text = "\n".join(" | ".join(cell if cell is not None else "" for cell in row) for row in value) if table else value
                blocks.append(Block(text=text, kind=kind, location=location, table=table))
    rendered_pdf.close()
    return blocks, issues


def image_blocks(data: bytes, source_type: str):
    from PIL import Image, ImageOps, ImageSequence

    try:
        image = Image.open(BytesIO(data))
        image.verify()
        image = Image.open(BytesIO(data))
        formats = {"JPEG": {"jpg", "jpeg"}, "TIFF": {"tif", "tiff"},
                   "PNG": {"png"}, "BMP": {"bmp"}, "WEBP": {"webp"}}
        if source_type not in formats.get(image.format, set()):
            raise ExtractionError("file extension does not match image content")
        frames = ImageSequence.Iterator(image)
        blocks, issues = [], []
        for page_number, frame in enumerate(frames, 1):
            if frame.width * frame.height > 50_000_000:
                raise ExtractionError("image dimensions exceed the 50 megapixel limit")
            image = ImageOps.exif_transpose(frame).convert("RGB")
            recognized = recognize_image(image)
            if not recognized:
                issues.append(Issue(code="OCR_NO_TEXT", message="この画像から文字を認識できませんでした。画像を確認してください。", page_number=page_number))
                continue
            text = "\n".join(line for line, _ in recognized)
            location = Location(page_number=page_number, block_index=len(blocks)+1,
                                locator=f"Image page {page_number}")
            blocks.append(Block(text=text, kind="text", location=location))
            if any(score is not None and score < 0.55 for _, score in recognized):
                issues.append(Issue(code="OCR_LOW_CONFIDENCE", message="OCR 識別精度が低い箇所があります。原文画像と照合してください。", page_number=page_number))
        return blocks, issues
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError(f"image extraction failed: {type(exc).__name__}: {exc}") from exc


def html_blocks(data: bytes, options: ExtractionOptions):
    from bs4 import BeautifulSoup, Tag

    if not options.html_selector:
        raise ExtractionError("HTML requires an explicit html_selector; whole-page guessing is disabled")
    soup = BeautifulSoup(data, "html.parser")
    roots = soup.select(options.html_selector)
    if len(roots) != 1:
        raise ExtractionError("html_selector must match exactly one element")
    root = roots[0]
    blocks, issues = [], []
    active = options.start_heading is None
    started = active
    ended = options.end_heading is None
    ordinal = 0
    for el in root.descendants:
        if not isinstance(el, Tag) or el.name not in {"p", "li", "h1", "h2", "h3", "h4", "h5", "h6", "table"}:
            continue
        if el.find_parent(["table", "script", "style", "nav", "aside"]):
            continue
        if el.name == "p" and el.find_parent("li"):
            continue
        if el.name == "li" and el.find_parent("li"):
            continue
        ordinal += 1
        value = el.get_text("\n", strip=True)
        is_heading = el.name.startswith("h")
        if is_heading and options.start_heading == value:
            if started:
                raise ExtractionError("start_heading is not unique")
            active = started = True
        if active and is_heading and options.end_heading == value:
            ended = True
            break
        if not active or not value:
            continue
        location = Location(page_number=None, block_index=len(blocks)+1,
                            locator=f"{options.html_selector}; semantic element {ordinal}; {el.name}")
        table = None
        if el.name == "table":
            rows = []
            spans = False
            for row in el.find_all("tr"):
                cells = row.find_all(["td", "th"], recursive=False)
                rows.append([cell.get_text("\n", strip=True) for cell in cells])
                spans |= any(cell.get("rowspan", "1") != "1" or cell.get("colspan", "1") != "1" for cell in cells)
            table = Table(rows=rows, location=location)
            if spans:
                issues.append(Issue(code="HTML_TABLE_SPANS", message="HTML 表に結合セルがあります。見出しの継承は推測せず、原文を確認してください。"))
        # Split physical lines so a paragraph with title + article can be recognized.
        for line in ([value] if table else value.splitlines()):
            if line.strip():
                loc = location.model_copy(update={"block_index": len(blocks)+1})
                blocks.append(Block(text=line, kind="table" if table else "heading" if is_heading else "text", location=loc, table=table))
    if not started or not ended:
        raise ExtractionError("configured start_heading or end_heading was not found exactly")
    return blocks, issues


def markdown_blocks(data: bytes):
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ExtractionError("Markdown must be UTF-8 encoded") from exc
    blocks = []
    heading_stack = []
    for line_number, raw_line in enumerate(text.splitlines(), 1):
        line = raw_line.strip()
        if not line:
            continue
        heading = re.match(r"^(#{1,6})\s+(.+?)\s*#*\s*$", line)
        if heading:
            level = len(heading.group(1))
            title = heading.group(2).strip()
            heading_stack = heading_stack[:level - 1]
            heading_stack.append(title)
            content, kind = title, "heading"
        else:
            content, kind = line, "text"
        location = Location(page_number=None, block_index=len(blocks) + 1,
                            locator=f"Markdown line {line_number}")
        blocks.append(Block(text=content, kind=kind, location=location))
    if not blocks:
        raise ExtractionError("no text extracted from Markdown")
    return blocks, []


def docx_blocks(data: bytes):
    from docx import Document
    from docx.table import Table as WordTable

    doc = Document(BytesIO(data))
    blocks, issues = [], []
    for n, element in enumerate(doc.iter_inner_content(), 1):
        loc = Location(page_number=None, block_index=n, locator=f"DOCX body element {n}")
        if isinstance(element, WordTable):
            rows = [[cell.text for cell in row.cells] for row in element.rows]
            table = Table(rows=rows, location=loc)
            blocks.append(Block(text="\n".join(" | ".join(row) for row in rows), kind="table", location=loc, table=table))
            if element._tbl.xpath(".//w:gridSpan|.//w:vMerge|.//w:tbl/w:tr/w:tc/w:tbl"):
                issues.append(Issue(code="DOCX_TABLE_COMPLEX", message="Word 表に結合セルまたは入れ子があります。セルの継承を推測せず、原文を確認してください。"))
        elif element.text.strip():
            for line in element.text.splitlines():
                if line.strip():
                    blocks.append(Block(text=line, kind="text", location=loc, table=None))
    if doc.inline_shapes or doc.element.xpath(".//w:txbxContent|.//w:del|.//w:ins"):
        issues.append(Issue(code="DOCX_NON_BODY_CONTENT", message="画像、テキストボックス、変更履歴を人が確認してください。"))
    # Header/footer data are deliberately not silently treated as body clauses.
    if any(p.text.strip() for s in doc.sections for part in (s.header, s.footer) for p in part.paragraphs):
        issues.append(Issue(code="DOCX_HEADER_FOOTER", message="ヘッダーとフッターは条文として取り込んでいません。規程情報が含まれているか確認してください。"))
    return blocks, issues


def extract(data: bytes, source_type: str, options: ExtractionOptions):
    try:
        if source_type == "pdf":
            blocks, issues = pdf_blocks(data)
        elif source_type == "html":
            blocks, issues = html_blocks(data, options)
        elif source_type == "docx":
            blocks, issues = docx_blocks(data)
        elif source_type == "md":
            blocks, issues = markdown_blocks(data)
        elif source_type in OCR_IMAGE_TYPES:
            blocks, issues = image_blocks(data, source_type)
        else:
            raise ExtractionError("unsupported explicit source type")
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError(f"{source_type} extraction failed: {type(exc).__name__}: {exc}") from exc
    if not blocks or not any(b.text.strip() for b in blocks):
        raise ExtractionError("no text extracted; image-only files require OCR")
    return blocks, issues


def chunk_blocks(blocks: list[Block], snapshot_id: str):
    chunks, current, heading_path = [], None, []
    pending_title = None

    def start(kind, label=None):
        nonlocal current
        current = Chunk(chunk_id=f"{snapshot_id}:{len(chunks)+1}", snapshot_id=snapshot_id,
                        sequence=len(chunks)+1, chunk_type=kind, heading_path=list(heading_path),
                        article_label=label, text="", locations=[], tables=[], references=[], issues=[])
        chunks.append(current)

    def append(block):
        current.text += ("\n" if current.text else "") + block.text
        current.locations.append(block.location)
        if block.table:
            current.tables.append(block.table)

    for block in blocks:
        text = block.text.strip()
        article = ARTICLE.match(text)
        table_label = TABLE_LABEL.match(text)
        if CHAPTER.match(text):
            if pending_title:
                if current is None:
                    start("context")
                append(pending_title)
                pending_title = None
            heading_path = [text]
            start("context")
        elif table_label:
            start("table", table_label.group(1))
        elif text in ("附則", "付則"):
            start("appendix", text)
        elif re.match(r"^[0-9]{4}年[0-9]+月[0-9]+日", text):
            if current is None or current.chunk_type != "context":
                start("context")
        elif article:
            start("article", article.group(1))
            if pending_title:
                current.heading_path.append(pending_title.text)
                append(pending_title)
                pending_title = None
        elif re.fullmatch(r"[（(][^）)]+[）)]", text) and block.kind != "table":
            if pending_title:
                if current is None:
                    start("context")
                append(pending_title)
            pending_title = block
            continue
        elif current is None:
            start("context")
        if pending_title and not article:
            append(pending_title)
            pending_title = None
        append(block)
    if pending_title:
        if current is None:
            start("context")
        append(pending_title)

    targets = {}
    for chunk in chunks:
        if chunk.article_label:
            targets.setdefault(chunk.article_label, []).append(chunk.chunk_id)
    for chunk in chunks:
        for label in dict.fromkeys(REFERENCE.findall(chunk.text)):
            if label == chunk.article_label:
                continue
            ids = targets.get(label, [])
            chunk.references.append(Reference(label=label, target_chunk_ids=ids,
                                    status="resolved" if len(ids) == 1 else "ambiguous" if ids else "unresolved"))
        if PLACEHOLDER.search(chunk.text):
            chunk.issues.append(Issue(code="UNFILLED_PLACEHOLDER", message="原文に未記入のプレースホルダーがあります。金額や日付として扱えません。", chunk_id=chunk.chunk_id))
        if any(r.status != "resolved" for r in chunk.references):
            chunk.issues.append(Issue(code="UNRESOLVED_REFERENCE", message="条文の参照先を解決できませんでした。近似する番号や文脈からは推測しません。", chunk_id=chunk.chunk_id))
    if not any(c.chunk_type == "article" for c in chunks):
        # Uploaded files may be plain instructions or scanned notes, not formal policies.
        for chunk in chunks:
            chunk.chunk_type = "context"
    return chunks
